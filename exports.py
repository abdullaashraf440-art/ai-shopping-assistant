"""Export helpers: CSV / styled Excel bytes and printable HTML receipts."""
from __future__ import annotations

import html
import io

import pandas as pd

from config import APP_NAME


def to_csv(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly


def to_xlsx(sheets: dict[str, pd.DataFrame]) -> bytes:
    """Write one or more DataFrames to a styled workbook (header fill, widths, freeze, filter)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="15803D")
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            if df.shape[1]:
                ws.auto_filter.ref = ws.dimensions
            sample = df.head(500)
            for i, col in enumerate(df.columns, 1):
                longest = max([len(str(col))] + [len(str(v)) for v in sample[col].tolist()])
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
    return buf.getvalue()


def receipt_html(order: dict) -> bytes:
    e = html.escape
    rows = "".join(
        f"<tr><td>{e(str(i['name']))}</td><td class=n>{i['quantity']}</td>"
        f"<td class=n>${i['unit_price']:.2f}</td><td class=n>${i['line_total']:.2f}</td></tr>"
        for i in order["items"])
    who = f"<p>Customer: <b>{e(order['customer_name'])}</b></p>" if order.get("customer_name") else ""
    page = f"""<!doctype html><html lang="en"><meta charset="utf-8">
<title>Receipt #{order['order_id']}</title>
<style>
body{{font:15px/1.5 system-ui,sans-serif;max-width:560px;margin:40px auto;padding:0 20px;color:#14221a}}
h1{{font-size:22px;margin:0}} .muted{{color:#5b6b61}} table{{width:100%;border-collapse:collapse;margin-top:16px}}
th,td{{padding:8px 6px;border-bottom:1px solid #e3e8e4;text-align:left}} .n{{text-align:right}}
tfoot td{{font-weight:700;border-bottom:none;font-size:17px}}
.badge{{display:inline-block;background:#dcfce7;color:#166534;padding:2px 10px;border-radius:99px;font-size:13px}}
@media print{{body{{margin:0}}}}
</style>
<h1>🛒 {e(APP_NAME)}</h1>
<p class=muted>Order #{order['order_id']} · {e(str(order['ordered_at']))} UTC</p>
<p><span class=badge>Confirmed</span> Estimated delivery in 3–5 business days</p>{who}
<table><thead><tr><th>Item</th><th class=n>Qty</th><th class=n>Price</th><th class=n>Total</th></tr></thead>
<tbody>{rows}</tbody>
<tfoot><tr><td colspan=3>Total</td><td class=n>${order['total']:.2f}</td></tr></tfoot></table>
<p class=muted>Thank you for shopping with us!</p></html>"""
    return page.encode("utf-8")
